from docx import Document
from pptx import Presentation
from pptx.util import Inches
from PIL import Image
from pathlib import Path
p=Path('/tmp/fuke-fixtures'); p.mkdir(exist_ok=True); Image.new('RGB',(120,70),(25,145,225)).save(p/'sample.png')
d=Document(); d.add_heading('DOCX 标题测试',0); d.add_heading('章节标题',1); d.add_paragraph('DOCX 正文内容'); d.add_paragraph('列表条目',style='List Bullet'); t=d.add_table(rows=2,cols=2); t.cell(0,0).text='表格姓名'; t.cell(0,1).text='表格数量'; t.cell(1,0).text='测试用户'; t.cell(1,1).text='42'; d.add_picture(str(p/'sample.png'),width=Inches(1)); d.save(p/'sample.docx')
r=Presentation()
for n in range(1,4):
 s=r.slides.add_slide(r.slide_layouts[5]); s.shapes.title.text=f'PPTX 第{n}页标题'; box=s.shapes.add_textbox(Inches(1),Inches(2), Inches(6),Inches(1)); box.text=f'第{n}页正文内容'; s.shapes.add_picture(str(p/'sample.png'),Inches(1),Inches(3),width=Inches(1))
r.save(p/'sample.pptx')
(p/'sample.md').write_text('# Markdown 标题\n\n**粗体内容**\n\n- 列表项目\n\n| 名称 | 值 |\n|---|---|\n| 测试 | 42 |\n\n<img src="x" onerror="window.xss=1"><script>window.xss=2</script>\n[恶意链接](javascript:alert(1))\n',encoding='utf-8')
(p/'broken.docx').write_text('not an office file')
(p/'broken.doc').write_text('not a legacy Office file')
(p/'huge.txt').write_bytes(b'x'*(10*1024*1024+1))
(p/'limit.txt').write_bytes(b'x'*(10*1024*1024))
(p/'unsafe.exe').write_text('not executable')

# Real PDF with distinct text and colored shapes on each of three pages.
from reportlab.pdfgen import canvas
c=canvas.Canvas(str(p/'sample.pdf'))
for n in range(1,4):
 c.setFont('Helvetica',32); c.drawString(60,740,f'PDF PAGE {n} / 3')
 c.setFillColorRGB(n/4,.3,1-n/4); c.rect(60,450,300,n*60,fill=1)
 c.setFillColorRGB(0,0,0); c.drawString(60,100,f'Real upload test - page {n}'); c.showPage()
c.save()
(p/'broken.pdf').write_text('not a pdf')
